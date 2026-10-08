"""Convert a markdown string to .docx bytes using python-docx."""

import io
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _apply_inline(paragraph, text: str) -> None:
    """Parse bold/italic markers in text and add runs to paragraph."""
    # Split on **bold** and *italic* markers
    pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)")
    last = 0
    for m in pattern.finditer(text):
        if m.start() > last:
            paragraph.add_run(text[last:m.start()])
        if m.group(2):  # **bold**
            paragraph.add_run(m.group(2)).bold = True
        elif m.group(3):  # *italic*
            paragraph.add_run(m.group(3)).italic = True
        elif m.group(4):  # `code`
            run = paragraph.add_run(m.group(4))
            run.font.name = "Courier New"
        last = m.end()
    if last < len(text):
        paragraph.add_run(text[last:])


def markdown_to_docx_bytes(markdown_text: str, photo_bytes: bytes | None = None) -> bytes:
    """Convert markdown to a .docx document and return raw bytes."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.1)
        section.right_margin  = Inches(1.1)

    # Default body font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if photo_bytes:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run()
        run.add_picture(io.BytesIO(photo_bytes), width=Inches(1.3))

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        # Headings
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)

        # Horizontal rule
        elif re.match(r"^[-*_]{3,}$", line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run("─" * 60)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

        # Unordered list
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _apply_inline(p, line[2:].strip())

        # Ordered list
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            _apply_inline(p, re.sub(r"^\d+\. ", "", line).strip())

        # Blank line → paragraph break
        elif line.strip() == "":
            doc.add_paragraph()

        # Normal paragraph
        else:
            p = doc.add_paragraph()
            _apply_inline(p, line.strip())

        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_filename(date_str: str, role: str, company: str, suffix: str) -> str:
    """Build a clean filename: YYYY-MM-DD_ShortRole_Company_suffix.docx"""
    def _slug(text: str, max_len: int) -> str:
        text = re.sub(r"[^\w\s-]", "", text.strip())
        text = re.sub(r"[\s-]+", "_", text)
        return text[:max_len]
    return f"{date_str}_{_slug(role, 20)}_{_slug(company, 15)}_{suffix}.docx"
